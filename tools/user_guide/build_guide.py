"""Builds docs/Litterbox_User_Guide.docx from guide_content.BLOCKS. Needs python-docx (pip install python-docx), which
the program itself doesn't use. Open the result in Word once and update its contents (right-click > Update Field).

    python tools/user_guide/build_guide.py docs/Litterbox_User_Guide.docx
"""

import re
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from guide_content import BLOCKS, VERSION

ACCENT = RGBColor(0x1F, 0x5F, 0x8B)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")
NUMBERS = re.compile(r"\b[A-Z]{2,6}-\d+(-\d+)?(_[A-Z]+)?")      # part and document numbers never break at a hyphen


def keep_together(text):
    return NUMBERS.sub(lambda m: m.group(0).replace("-", "‑"), text)


def runs(paragraph, text, size=None):
    for piece in INLINE.split(keep_together(text)):
        if not piece:
            continue
        if piece.startswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
        else:
            run = paragraph.add_run(piece)
        if size:
            run.font.size = size


def shade(cell_or_par, fill):
    target = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), fill)
    target.append(element)


def left_border(paragraph, color="1F5F8B"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for key, value in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", color)):
        left.set(qn(f"w:{key}"), value)
    borders.append(left)
    ppr.append(borders)


def field(paragraph, instruction):
    """A Word field (page number, table of contents)."""
    run = paragraph.add_run()
    begin, text, separate, end = (OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar"),
                                  OxmlElement("w:fldChar"))
    begin.set(qn("w:fldCharType"), "begin")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate.set(qn("w:fldCharType"), "separate")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(separate)
    placeholder = paragraph.add_run("Right-click and choose Update Field to fill in the contents.")
    placeholder.italic = True
    paragraph.add_run()._r.append(end)


def styles(document):
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(11), INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Heading 1", 18, 18, 6), ("Heading 2", 13.5, 14, 4), ("Heading 3", 11.5, 10, 2)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.bold = "Calibri", Pt(size), True
        style.font.color.rgb = ACCENT if name != "Heading 3" else INK
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        document.styles[name].paragraph_format.space_after = Pt(3)


def table(document, headings, rows, widths):
    grid = document.add_table(rows=1, cols=len(headings))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for cell, text, width in zip(grid.rows[0].cells, headings, widths):
        cell.width = Inches(width)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10), RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "1F5F8B")
    for number, row in enumerate(rows):
        cells = grid.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            runs(cell.paragraphs[0], text, Pt(10))
            if number % 2:
                shade(cell, "EEF3F7")
    for column, width in zip(grid.columns, widths):
        column.width = Inches(width)
    header = grid.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def new_numbering(document):
    """A fresh numbering for one numbered list, so each list starts at 1."""
    numbering = document.part.numbering_part.numbering_definitions._numbering
    style_num = document.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


def numbered(paragraph, num_id):
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def cover(document):
    for _ in range(6):
        document.add_paragraph()
    title = document.add_paragraph()
    run = title.add_run("Litterbox PDM")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(36), ACCENT
    subtitle = document.add_paragraph()
    run = subtitle.add_run("User Guide")
    run.font.size, run.font.color.rgb = Pt(22), INK
    rule = document.add_paragraph()
    ppr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "12"), ("space", "1"), ("color", "1F5F8B")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    ppr.append(borders)
    about = document.add_paragraph()
    runs(about, "Check-in, check-out, revisions, releases and BOMs for small SolidWorks teams, on Git and Git LFS.")
    about.runs[0].font.color.rgb = MUTED
    version = document.add_paragraph()
    run = version.add_run(f"For Litterbox {VERSION}  ·  {date.today():%B %Y}")
    run.font.color.rgb = MUTED
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    heading = document.add_paragraph()                      # not a Heading, so it isn't listed in itself
    run = heading.add_run("Contents")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(18), ACCENT
    heading.paragraph_format.space_after = Pt(8)
    field(document.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(document):
    section = document.sections[0]
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Litterbox User Guide  ·  page ")
    run.font.size, run.font.color.rgb = Pt(9), MUTED
    field_paragraph = paragraph
    run = field_paragraph.add_run()
    begin, text, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text.set(qn("xml:space"), "preserve")
    text.text = "PAGE"
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(text)
    run._r.append(end)
    run.font.size = Pt(9)
    section.different_first_page_header_footer = True


def build(path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.top_margin, section.bottom_margin = Inches(0.9), Inches(0.9)
    styles(document)
    cover(document)
    footer(document)
    first = True
    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            if not first:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first = False
            document.add_paragraph(block[1], style="Heading 1")
        elif kind in ("h2", "h3"):
            document.add_paragraph(block[1], style="Heading 2" if kind == "h2" else "Heading 3")
        elif kind == "p":
            runs(document.add_paragraph(), block[1])
        elif kind in ("ul", "ol"):
            restart = new_numbering(document) if kind == "ol" else None
            for item in block[1]:
                paragraph = document.add_paragraph(style="List Bullet" if kind == "ul" else "List Number")
                if restart is not None:
                    numbered(paragraph, restart)
                runs(paragraph, item)
        elif kind == "code":
            for line in block[1]:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.3)
                paragraph.paragraph_format.space_after = Pt(1)
                run = paragraph.add_run(line)
                run.font.name, run.font.size = "Consolas", Pt(9.5)
            document.paragraphs[-1].paragraph_format.space_after = Pt(8)
        elif kind == "table":
            table(document, block[1], block[2], block[3])
        elif kind == "tip":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.1)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(10)
            shade(paragraph, "EEF3F7")
            left_border(paragraph)
            runs(paragraph, block[1])
        else:
            raise ValueError(kind)
    document.core_properties.title = "Litterbox PDM User Guide"
    document.core_properties.subject = f"Litterbox {VERSION}"
    document.save(path)


if __name__ == "__main__":
    build(sys.argv[1])
    print("written", sys.argv[1])
